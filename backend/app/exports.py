from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import cm
from reportlab.platypus import (
    SimpleDocTemplate,
    Paragraph,
    Spacer,
    Table,
    TableStyle,
)


def export_invoices_to_excel(invoices: list) -> BytesIO:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Factures"

    headers = [
        "N° facture",
        "Date facture",
        "Date échéance",
        "Fournisseur",
        "Client",
        "Total HT",
        "TVA",
        "Total TTC",
        "Devise",
        "Fichier",
    ]

    sheet.append(headers)

    for item in invoices:
        invoice = item.get("invoice", {})

        sheet.append([
            invoice.get("invoice_number") or "",
            invoice.get("issue_date") or "",
            invoice.get("due_date") or "",
            invoice.get("supplier") or "",
            invoice.get("customer") or "",
            invoice.get("total_ht") or 0,
            invoice.get("tax") or 0,
            invoice.get("total_ttc") or 0,
            invoice.get("currency") or "",
            item.get("filename") or "",
        ])

    header_fill = PatternFill(
        fill_type="solid",
        fgColor="1D4ED8",
    )

    for cell in sheet[1]:
        cell.font = Font(
            bold=True,
            color="FFFFFF",
        )
        cell.fill = header_fill
        cell.alignment = Alignment(
            horizontal="center",
        )

    for column in sheet.columns:
        max_length = max(
            len(str(cell.value or ""))
            for cell in column
        )
        sheet.column_dimensions[
            column[0].column_letter
        ].width = min(max_length + 3, 35)

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions

    output = BytesIO()
    workbook.save(output)
    output.seek(0)

    return output


def export_invoices_to_pdf(invoices: list) -> BytesIO:
    output = BytesIO()

    document = SimpleDocTemplate(
        output,
        pagesize=landscape(A4),
        rightMargin=1 * cm,
        leftMargin=1 * cm,
        topMargin=1.5 * cm,
        bottomMargin=1.5 * cm,
    )

    styles = getSampleStyleSheet()

    title_style = ParagraphStyle(
        "InvoiceTitle",
        parent=styles["Title"],
        alignment=TA_CENTER,
        textColor=colors.HexColor("#1D4ED8"),
    )

    elements = [
        Paragraph(
            "Rapport des factures",
            title_style,
        ),
        Spacer(1, 0.5 * cm),
    ]

    headers = [
        "N° facture",
        "Date",
        "Fournisseur",
        "Client",
        "Total HT",
        "TVA",
        "Total TTC",
        "Devise",
    ]

    data = [headers]

    for item in invoices:
        invoice = item.get("invoice", {})

        data.append([
            str(invoice.get("invoice_number") or "-"),
            str(invoice.get("issue_date") or "-"),
            str(invoice.get("supplier") or "-"),
            str(invoice.get("customer") or "-"),
            str(invoice.get("total_ht") or 0),
            str(invoice.get("tax") or 0),
            str(invoice.get("total_ttc") or 0),
            str(invoice.get("currency") or "-"),
        ])

    table = Table(
        data,
        repeatRows=1,
        colWidths=[
            3.0 * cm,
            2.2 * cm,
            4.0 * cm,
            4.0 * cm,
            2.2 * cm,
            2.0 * cm,
            2.4 * cm,
            1.8 * cm,
        ],
    )

    table.setStyle(TableStyle([
        (
            "BACKGROUND",
            (0, 0),
            (-1, 0),
            colors.HexColor("#1D4ED8"),
        ),
        (
            "TEXTCOLOR",
            (0, 0),
            (-1, 0),
            colors.white,
        ),
        (
            "FONTNAME",
            (0, 0),
            (-1, 0),
            "Helvetica-Bold",
        ),
        (
            "FONTSIZE",
            (0, 0),
            (-1, -1),
            8,
        ),
        (
            "GRID",
            (0, 0),
            (-1, -1),
            0.5,
            colors.HexColor("#D1D5DB"),
        ),
        (
            "VALIGN",
            (0, 0),
            (-1, -1),
            "MIDDLE",
        ),
        (
            "LEFTPADDING",
            (0, 0),
            (-1, -1),
            5,
        ),
        (
            "RIGHTPADDING",
            (0, 0),
            (-1, -1),
            5,
        ),
        (
            "TOPPADDING",
            (0, 0),
            (-1, -1),
            7,
        ),
        (
            "BOTTOMPADDING",
            (0, 0),
            (-1, -1),
            7,
        ),
    ]))

    elements.append(table)

    document.build(elements)
    output.seek(0)

    return output